"""Update ticket workbooks after executing a generated schedule.

Reads the original input ticket files and a schedule workbook, subtracts the
one ticket consumed by each scheduled battle, and saves updated _票.xlsx files.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .config import (
    TARGETS,
    WILDCARD,
    CHARACTER_COL,
    CHAR_SHEET,
    SCHEDULE_SHEET,
    TICKET_SUFFIX,
)


@dataclass
class TicketBook:
    member: str
    path: Path
    workbook: object
    sheet: Worksheet
    columns: Dict[str, int]
    rows_by_character: Dict[str, int]


def _cell_to_int(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, str) and value.strip() == "":
        return 0
    try:
        return max(0, int(float(value)))
    except (TypeError, ValueError):
        return 0


def _stripped(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_map(ws: Worksheet) -> Dict[str, int]:
    return {
        _stripped(cell.value): cell.column
        for cell in ws[1]
        if _stripped(cell.value)
    }


def _load_ticket_books(input_dir: Path) -> Dict[str, TicketBook]:
    books: Dict[str, TicketBook] = {}
    required_cols = {CHARACTER_COL, *TARGETS, WILDCARD}

    for path in sorted(input_dir.glob(f"*{TICKET_SUFFIX}.xlsx")):
        member = path.stem[: -len(TICKET_SUFFIX)]
        wb = load_workbook(path)
        if CHAR_SHEET not in wb.sheetnames:
            raise ValueError(f"{path.name} missing sheet '{CHAR_SHEET}'")
        ws = wb[CHAR_SHEET]
        columns = _header_map(ws)
        missing = required_cols - set(columns)
        if missing:
            raise ValueError(f"{path.name} missing columns: {sorted(missing)}")

        char_col = columns[CHARACTER_COL]
        rows_by_character: Dict[str, int] = {}
        for row in range(2, ws.max_row + 1):
            character = _stripped(ws.cell(row=row, column=char_col).value)
            if not character:
                continue
            if character in rows_by_character:
                raise ValueError(f"Duplicate character '{character}' in {path.name}")
            rows_by_character[character] = row

        if rows_by_character:
            books[member] = TicketBook(
                member=member,
                path=path,
                workbook=wb,
                sheet=ws,
                columns=columns,
                rows_by_character=rows_by_character,
            )

    if not books:
        raise ValueError(f"No ticket files found in {input_dir}")
    return books


def _iter_schedule_rows(schedule_path: Path) -> Iterable[Tuple[int, str, str, str, str]]:
    wb = load_workbook(schedule_path, data_only=True)
    if SCHEDULE_SHEET not in wb.sheetnames:
        raise ValueError(f"{schedule_path.name} missing sheet '{SCHEDULE_SHEET}'")
    ws = wb[SCHEDULE_SHEET]
    columns = _header_map(ws)
    required = {"order", "target", "ticket_kind", "ticket_source"}
    missing = required - set(columns)
    if missing:
        raise ValueError(f"{schedule_path.name} missing columns: {sorted(missing)}")

    for row in range(2, ws.max_row + 1):
        target = _stripped(ws.cell(row=row, column=columns["target"]).value)
        if not target:
            continue
        raw_order = ws.cell(row=row, column=columns["order"]).value
        try:
            order = int(raw_order)
        except (TypeError, ValueError):
            order = row - 1
        ticket_kind = _stripped(ws.cell(row=row, column=columns["ticket_kind"]).value)
        ticket_source = _stripped(ws.cell(row=row, column=columns["ticket_source"]).value)
        if ":" not in ticket_source:
            raise ValueError(f"battle #{order} has malformed ticket_source '{ticket_source}'")
        member, character = (part.strip() for part in ticket_source.split(":", 1))
        yield order, target, ticket_kind, member, character


def update_ticket_files(
    input_dir: str | Path,
    schedule_path: str | Path,
    output_dir: str | Path | None = None,
    in_place: bool = False,
) -> dict:
    """Apply schedule ticket consumption and save updated ticket workbooks.

    If `in_place` is False, updated files are written to `output_dir` using the
    original filenames. If `in_place` is True, original ticket files are
    overwritten.
    """
    input_dir = Path(input_dir)
    schedule_path = Path(schedule_path)
    if in_place:
        save_dir = input_dir
    else:
        save_dir = Path(output_dir) if output_dir is not None else Path("remaining_tickets")
        save_dir.mkdir(parents=True, exist_ok=True)

    books = _load_ticket_books(input_dir)
    summary: dict = {
        "dedicated": 0,
        "wildcard": 0,
        "by_member": {m: {"dedicated": 0, "wildcard": 0} for m in books},
        "saved_files": [],
    }
    errors: list[str] = []

    for order, target, kind, member, character in _iter_schedule_rows(schedule_path):
        book = books.get(member)
        if book is None:
            errors.append(f"battle #{order} {target}: no ticket workbook for member '{member}'")
            continue
        row = book.rows_by_character.get(character)
        if row is None:
            errors.append(f"battle #{order} {target}: character '{character}' not found in {book.path.name}")
            continue

        if kind == WILDCARD:
            if target == "双生":
                errors.append(f"battle #{order} {target}: 双生 cannot consume '{WILDCARD}'")
                continue
            column_name = WILDCARD
            kind_key = "wildcard"
        elif kind == target:
            if target not in TARGETS:
                errors.append(f"battle #{order} {target}: unknown target")
                continue
            column_name = target
            kind_key = "dedicated"
        else:
            errors.append(
                f"battle #{order} {target}: ticket_kind '{kind}' must be '{target}' or '{WILDCARD}'"
            )
            continue

        col = book.columns[column_name]
        cell = book.sheet.cell(row=row, column=col)
        current = _cell_to_int(cell.value)
        if current <= 0:
            errors.append(
                f"battle #{order} {target}: {member}:{character} has no '{column_name}' tickets left"
            )
            continue
        cell.value = current - 1
        summary[kind_key] += 1
        summary["by_member"].setdefault(member, {"dedicated": 0, "wildcard": 0})[kind_key] += 1

    if errors:
        raise ValueError("Ticket update failed:\n" + "\n".join(f"- {e}" for e in errors))

    for member, book in books.items():
        save_path = book.path if in_place else save_dir / book.path.name
        book.workbook.save(save_path)
        summary["saved_files"].append(str(save_path))

    return summary


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Subtract scheduled ticket usage from _票.xlsx workbooks."
    )
    parser.add_argument("--input-dir", default="input", help="Directory containing <member>_票.xlsx files.")
    parser.add_argument("--schedule", default="schedule.xlsx", help="Schedule workbook produced by the optimizer.")
    parser.add_argument(
        "--out-dir",
        default="remaining_tickets",
        help="Where to write updated ticket files when not using --in-place.",
    )
    parser.add_argument(
        "--in-place",
        action="store_true",
        help="Overwrite original _票.xlsx files in --input-dir.",
    )
    args = parser.parse_args()

    summary = update_ticket_files(
        args.input_dir,
        args.schedule,
        output_dir=args.out_dir,
        in_place=args.in_place,
    )
    total = summary["dedicated"] + summary["wildcard"]
    print(f"Applied ticket usage: dedicated={summary['dedicated']}, wildcard={summary['wildcard']}, total={total}")
    for member, counts in summary["by_member"].items():
        used = counts["dedicated"] + counts["wildcard"]
        if used:
            print(f"  {member}: dedicated={counts['dedicated']}, wildcard={counts['wildcard']}, total={used}")
    print("Saved updated ticket files:")
    for path in summary["saved_files"]:
        print(f"  {path}")


if __name__ == "__main__":
    main()
