"""Regression checks for updating ticket workbooks after a schedule.

Run:
    python -m tests.test_update_tickets
"""
from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.config import (
    CHAR_SHEET,
    SCHEDULE_SHEET,
    TARGETS,
    WILDCARD,
    CHARACTER_COL,
    TICKET_SUFFIX,
)
from src.update_tickets import update_ticket_files


def _write_ticket_file(path: Path, character: str, values: dict[str, int | str | None]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = CHAR_SHEET
    headers = [CHARACTER_COL] + list(TARGETS) + [WILDCARD]
    ws.append(headers)
    row = [character]
    for target in TARGETS:
        row.append(values.get(target, ""))
    row.append(values.get(WILDCARD, ""))
    ws.append(row)
    wb.save(path)


def _write_schedule(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SCHEDULE_SHEET
    ws.append(["order", "target", "ticket_kind", "ticket_source", "Alpha", "Bravo", "quests_completed"])
    ws.append([1, "狮蝎", "狮蝎", "Alpha:A1", "A1", "B1", 2])
    ws.append([2, "海龙", WILDCARD, "Alpha:A1", "A1", "B1", 1])
    ws.append([3, "海龙", "海龙", "Bravo:B1", "A1", "B1", 1])
    wb.save(path)


def _value(path: Path, column_name: str) -> object:
    wb = load_workbook(path)
    ws = wb[CHAR_SHEET]
    headers = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    return ws.cell(row=2, column=headers[column_name]).value


def main() -> None:
    root = Path(__file__).resolve().parent
    input_dir = root / "tmp_update_tickets_input"
    out_dir = root / "tmp_update_tickets_output"
    schedule_path = root / "tmp_update_tickets_schedule.xlsx"

    for p in (input_dir, out_dir):
        if p.exists():
            shutil.rmtree(p)
    if schedule_path.exists():
        schedule_path.unlink()
    input_dir.mkdir(parents=True)

    _write_ticket_file(
        input_dir / f"Alpha{TICKET_SUFFIX}.xlsx",
        "A1",
        {"狮蝎": 2, "海龙": "", WILDCARD: 1},
    )
    _write_ticket_file(
        input_dir / f"Bravo{TICKET_SUFFIX}.xlsx",
        "B1",
        {"海龙": 1, WILDCARD: None},
    )
    _write_schedule(schedule_path)

    summary = update_ticket_files(input_dir, schedule_path, output_dir=out_dir)
    assert summary["dedicated"] == 2
    assert summary["wildcard"] == 1

    alpha_out = out_dir / f"Alpha{TICKET_SUFFIX}.xlsx"
    bravo_out = out_dir / f"Bravo{TICKET_SUFFIX}.xlsx"
    assert _value(alpha_out, "狮蝎") == 1
    assert _value(alpha_out, WILDCARD) == 0
    assert _value(bravo_out, "海龙") == 0

    # Original input files are untouched in non-in-place mode.
    assert _value(input_dir / f"Alpha{TICKET_SUFFIX}.xlsx", "狮蝎") == 2
    assert _value(input_dir / f"Alpha{TICKET_SUFFIX}.xlsx", WILDCARD) == 1
    assert _value(input_dir / f"Bravo{TICKET_SUFFIX}.xlsx", "海龙") == 1

    print("[OK] schedule ticket consumption updates remaining ticket workbooks correctly.")


if __name__ == "__main__":
    main()
